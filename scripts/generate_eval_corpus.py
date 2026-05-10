from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT_DIR = ROOT / "eval_corpus"


DOCUMENTS: list[dict[str, Any]] = [
    {
        "file_name": "eval_reglament_client_requests_2025.docx",
        "title": "Регламент обработки клиентских запросов 2025",
        "doc_year": 2025,
        "doc_category": "регламент, клиентские запросы",
        "description": "Синтетический документ для проверки RAG по клиентским запросам.",
        "sections": [
            (
                "Прием запроса",
                [
                    "Клиентский запрос регистрируется в журнале обращений в течение 15 минут после поступления.",
                    "Оператор обязан присвоить обращению уникальный номер и указать канал поступления: портал, электронная почта или телефон.",
                    "Если запрос содержит персональные данные, оператор применяет режим ограниченного доступа до завершения идентификации заявителя.",
                ],
            ),
            (
                "Идентификация заявителя",
                [
                    "Идентификация заявителя выполняется по двум независимым признакам: номеру договора и одноразовому коду подтверждения.",
                    "Если один из признаков не подтвержден, ответ по существу запроса не предоставляется.",
                    "Повторная идентификация требуется при изменении канала связи или при запросе сведений кредитной истории.",
                ],
            ),
            (
                "Сроки подготовки ответа",
                [
                    "Стандартный срок подготовки ответа на клиентский запрос составляет три рабочих дня.",
                    "Сложный запрос может быть продлен до десяти рабочих дней по решению руководителя группы.",
                    "Факт продления и причина продления фиксируются в журнале обращений.",
                ],
            ),
        ],
    },
    {
        "file_name": "eval_knowledge_base_ingestion_2025.docx",
        "title": "Инструкция по загрузке документов в базу знаний",
        "doc_year": 2025,
        "doc_category": "база знаний, загрузка документов",
        "description": "Синтетический документ для проверки ingestion и метаданных.",
        "sections": [
            (
                "Поддерживаемые форматы",
                [
                    "Система поддерживает загрузку документов только в форматах PDF и DOCX.",
                    "Файлы других форматов отклоняются до этапа извлечения текста.",
                    "Пустой файл не индексируется и возвращает ошибку Empty file.",
                ],
            ),
            (
                "Индексирование",
                [
                    "После загрузки документ разбивается на текстовые чанки размером 1000 символов с перекрытием 150 символов.",
                    "Для каждого чанка рассчитывается embedding модели nomic-embed-text размерностью 768.",
                    "Дубликаты документов определяются по SHA-256 хэшу исходного файла.",
                ],
            ),
            (
                "Метаданные",
                [
                    "Каждый документ может иметь год документа, категорию, описание и признак активности.",
                    "Неактивные документы скрываются из поиска по умолчанию и не должны использоваться в ответах пользователю.",
                    "Категории нормализуются и хранятся как набор уникальных значений.",
                ],
            ),
        ],
    },
    {
        "file_name": "eval_security_audit_2025.docx",
        "title": "Политика безопасности и аудита запросов",
        "doc_year": 2025,
        "doc_category": "безопасность, аудит",
        "description": "Синтетический документ для проверки безопасности, RBAC и аудита.",
        "sections": [
            (
                "Роли доступа",
                [
                    "Роль user разрешает задавать вопросы чат-боту и отправлять обратную связь.",
                    "Роль admin разрешает загружать документы, изменять метаданные, активировать и удалять документы.",
                    "Если у пользователя нет подходящей роли, система возвращает ошибку Insufficient role.",
                ],
            ),
            (
                "Аудит действий",
                [
                    "Каждый запрос к чат-боту записывается в audit_log с действием ask и уникальным request_id.",
                    "Операции загрузки документов записываются в audit_log с действием upload и статусом indexed или already_indexed.",
                    "Комментарии обратной связи записываются после редактирования чувствительных данных.",
                ],
            ),
            (
                "Редактирование чувствительных данных",
                [
                    "Перед записью в журнал система маскирует email, телефонные номера и номера счетов.",
                    "В открытом ответе пользователю не должны появляться служебные секреты, JWT-токены или пароли.",
                    "Журнал аудита хранит только редактированную полезную нагрузку.",
                ],
            ),
        ],
    },
    {
        "file_name": "eval_rag_answer_policy_2025.docx",
        "title": "Политика формирования ответов RAG",
        "doc_year": 2025,
        "doc_category": "rag, качество ответов",
        "description": "Синтетический документ для проверки groundedness ответов.",
        "sections": [
            (
                "Использование контекста",
                [
                    "Ответ формируется только на основе найденных фрагментов контекста.",
                    "Если в контексте нет ответа, ассистент должен написать: Информация недостаточна для ответа.",
                    "Модель не должна добавлять предположения, внешние факты или неподтвержденные рекомендации.",
                ],
            ),
            (
                "Цитирование источников",
                [
                    "При использовании фрагментов контекста ответ должен содержать ссылки вида [source1], [source2].",
                    "Номер ссылки должен соответствовать реально найденному источнику в списке retrieved chunks.",
                    "Ответ без ссылок считается неполным для регламентных вопросов.",
                ],
            ),
            (
                "Проверка качества",
                [
                    "Метрика groundedness оценивает, связан ли ответ с найденным контекстом.",
                    "Метрика retrieval_hit_rate показывает долю вопросов, для которых найден хотя бы один чанк.",
                    "Метрика answer_context_overlap_avg измеряет среднее пересечение значимых токенов ответа с контекстом.",
                ],
            ),
        ],
    },
]


QA_ROWS: list[dict[str, Any]] = [
    {
        "id": "eval_q01",
        "question": "За какое время клиентский запрос регистрируется в журнале обращений?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "регламент",
        "expected_source_files": ["eval_reglament_client_requests_2025.docx"],
        "expected_answer_substrings": ["15 минут"],
    },
    {
        "id": "eval_q02",
        "question": "Какие два признака используются для идентификации заявителя?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "клиентские запросы",
        "expected_source_files": ["eval_reglament_client_requests_2025.docx"],
        "expected_answer_substrings": ["номер договора", "одноразовому коду"],
    },
    {
        "id": "eval_q03",
        "question": "Какой стандартный срок подготовки ответа на клиентский запрос?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_reglament_client_requests_2025.docx"],
        "expected_answer_substrings": ["три рабочих дня"],
    },
    {
        "id": "eval_q04",
        "question": "До какого срока может быть продлен сложный запрос?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_reglament_client_requests_2025.docx"],
        "expected_answer_substrings": ["десяти рабочих дней"],
    },
    {
        "id": "eval_q05",
        "question": "Какие форматы документов поддерживает загрузка в базу знаний?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "база знаний",
        "expected_source_files": ["eval_knowledge_base_ingestion_2025.docx"],
        "expected_answer_substrings": ["PDF", "DOCX"],
    },
    {
        "id": "eval_q06",
        "question": "Как система определяет дубликаты документов?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "загрузка документов",
        "expected_source_files": ["eval_knowledge_base_ingestion_2025.docx"],
        "expected_answer_substrings": ["SHA-256"],
    },
    {
        "id": "eval_q07",
        "question": "Какая размерность embedding используется для каждого чанка?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_knowledge_base_ingestion_2025.docx"],
        "expected_answer_substrings": ["768"],
    },
    {
        "id": "eval_q08",
        "question": "Что происходит с неактивными документами при поиске по умолчанию?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_knowledge_base_ingestion_2025.docx"],
        "expected_answer_substrings": ["скрываются из поиска"],
    },
    {
        "id": "eval_q09",
        "question": "Какие действия разрешены роли user?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "безопасность",
        "expected_source_files": ["eval_security_audit_2025.docx"],
        "expected_answer_substrings": ["задавать вопросы", "обратную связь"],
    },
    {
        "id": "eval_q10",
        "question": "Какие действия разрешены роли admin?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "аудит",
        "expected_source_files": ["eval_security_audit_2025.docx"],
        "expected_answer_substrings": ["загружать документы", "удалять документы"],
    },
    {
        "id": "eval_q11",
        "question": "Как записывается запрос к чат-боту в журнал аудита?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_security_audit_2025.docx"],
        "expected_answer_substrings": ["действием ask", "request_id"],
    },
    {
        "id": "eval_q12",
        "question": "Какие чувствительные данные маскируются перед записью в журнал?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_security_audit_2025.docx"],
        "expected_answer_substrings": ["email", "телефонные номера", "номера счетов"],
    },
    {
        "id": "eval_q13",
        "question": "На основе чего должен формироваться ответ RAG?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "rag",
        "expected_source_files": ["eval_rag_answer_policy_2025.docx"],
        "expected_answer_substrings": ["найденных фрагментов контекста"],
    },
    {
        "id": "eval_q14",
        "question": "Что должен написать ассистент, если в контексте нет ответа?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "doc_category": "качество ответов",
        "expected_source_files": ["eval_rag_answer_policy_2025.docx"],
        "expected_answer_substrings": ["Информация недостаточна для ответа"],
    },
    {
        "id": "eval_q15",
        "question": "Какой формат ссылок на источники должен использовать ответ?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_rag_answer_policy_2025.docx"],
        "expected_answer_substrings": ["[source1]", "[source2]"],
    },
    {
        "id": "eval_q16",
        "question": "Что показывает метрика retrieval_hit_rate?",
        "doc_year_from": 2025,
        "doc_year_to": 2025,
        "expected_source_files": ["eval_rag_answer_policy_2025.docx"],
        "expected_answer_substrings": ["долю вопросов"],
    },
]


def write_docx(path: Path, spec: dict[str, Any]) -> None:
    doc = Document()
    doc.add_heading(spec["title"], level=0)
    doc.add_paragraph(f"Год документа: {spec['doc_year']}")
    doc.add_paragraph(f"Категории: {spec['doc_category']}")
    doc.add_paragraph(spec["description"])

    for heading, paragraphs in spec["sections"]:
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)

    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR))
    args = parser.parse_args()

    out_dir = Path(args.out_dir)
    docs_dir = out_dir / "documents"
    docs_dir.mkdir(parents=True, exist_ok=True)

    manifest = []
    for spec in DOCUMENTS:
        doc_path = docs_dir / spec["file_name"]
        write_docx(doc_path, spec)
        manifest.append(
            {
                "file_name": spec["file_name"],
                "path": doc_path.relative_to(out_dir).as_posix(),
                "doc_year": spec["doc_year"],
                "doc_category": spec["doc_category"],
                "description": spec["description"],
                "is_active": True,
            }
        )

    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    with (out_dir / "qa_set.jsonl").open("w", encoding="utf-8", newline="\n") as f:
        for row in QA_ROWS:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    print(f"Wrote {len(manifest)} documents to {docs_dir}")
    print(f"Wrote QA set to {out_dir / 'qa_set.jsonl'}")


if __name__ == "__main__":
    main()
